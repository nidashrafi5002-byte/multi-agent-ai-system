from fastapi import APIRouter, UploadFile, File, HTTPException
from pydantic import BaseModel
import asyncio
import re
import io
from time import perf_counter
from agents.router import route_query
from pipelines.research_pipeline import run_research_pipeline
from pipelines.stock_pipeline import run_stock_pipeline, get_stock_chart
from pipelines.code_pipeline import run_code_pipeline
from pipelines.job_pipeline import run_job_pipeline
from pipelines.flight_pipeline import run_flight_pipeline
from pipelines.image_pipeline import generate_image
from pipelines.general_pipeline import run_general_pipeline
from workflows.graph_workflow import run_graph

router = APIRouter()

class ChatRequest(BaseModel):
    message: str
    pipeline: str = "auto"

@router.post("/api/chat")
async def chat(req: ChatRequest):
    # The agent pipelines perform blocking provider/database calls. Keep them
    # off FastAPI's event loop so health checks and other requests stay alive.
    result = await asyncio.to_thread(route_query, req.message)
    domain = result["domain"]
    started_at = perf_counter()
    if req.pipeline != "auto" and req.pipeline in {"research", "stock", "code", "job", "flight", "image", "general"}:
        domain = req.pipeline

    execution_log = [{"step": f"Orchestrator routed to {domain.title()} Agent", "duration_ms": round((perf_counter() - started_at) * 1000)}]

    def with_log(response):
        execution_log.append({"step": f"{domain.title()} pipeline completed", "duration_ms": round((perf_counter() - started_at) * 1000) - execution_log[0]["duration_ms"]})
        response["execution_log"] = execution_log
        return response
    
    if domain == "research":
        graph_result = await asyncio.to_thread(run_graph, req.message)
        return with_log({"domain": domain, "output": graph_result["written_report"], "score": graph_result["quality_score"]})
    elif domain == "stock":
        output = await asyncio.to_thread(run_stock_pipeline, req.message)
        symbol_map = {
            "google": "GOOG", "alphabet": "GOOG", "apple": "AAPL",
            "microsoft": "MSFT", "tesla": "TSLA", "amazon": "AMZN",
            "nvidia": "NVDA", "meta": "META"
        }
        query_lower = req.message.lower()
        symbol = next((value for key, value in symbol_map.items() if key in query_lower), None)
        if symbol is None:
            symbol_match = re.search(r"\b[A-Z]{1,5}(?:\.[A-Z]{1,3})?\b", req.message)
            symbol = symbol_match.group(0) if symbol_match else None

        chart_html = None
        if symbol:
            chart = await asyncio.to_thread(get_stock_chart, symbol)
            if chart:
                chart_html = chart.to_html(full_html=False, include_plotlyjs="cdn")

        return with_log({"domain": domain, "output": output, "stock_chart": chart_html})
    elif domain == "code":
        output = await asyncio.to_thread(run_code_pipeline, req.message)
        return with_log({"domain": domain, "output": output})
    elif domain == "job":
        output = await asyncio.to_thread(run_job_pipeline, req.message)
        return with_log({"domain": domain, "output": output})
    elif domain == "flight":
        report, flight_map, enriched = await asyncio.to_thread(run_flight_pipeline, req.message)
        map_html = flight_map.get_root().render() if flight_map else None
        return with_log({"domain": domain, "output": report, "enriched": enriched, "map_html": map_html})
    elif domain == "image":
        result = await asyncio.to_thread(generate_image, req.message)
        return with_log({
            "domain": domain,
            "output": f"Generated image for: {result['original_query']}",
            "image_url": result["image_url"],
            "enhanced_prompt": result["enhanced_prompt"]
        })
    else:
        output = await asyncio.to_thread(run_general_pipeline, req.message)
        return with_log({"domain": domain, "output": output})


@router.post("/api/analyze-file")
async def analyze_file(file: UploadFile = File(...)):
    """Analyze uploaded file (PDF, DOCX, or image) and return extracted content."""
    try:
        contents = await file.read()
        filename = file.filename or "uploaded_file"
        filename_lower = filename.lower()

        # Handle PDF files
        if filename_lower.endswith(".pdf"):
            try:
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(contents))
                text = ""
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                return {
                    "success": True,
                    "file_type": "pdf",
                    "filename": filename,
                    "text": text.strip(),
                    "pages": len(reader.pages)
                }
            except ImportError:
                raise HTTPException(status_code=500, detail="pypdf is required for PDF processing")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read PDF: {str(e)}")

        # Handle DOCX files
        elif filename_lower.endswith(".docx"):
            try:
                import docx
                doc = docx.Document(io.BytesIO(contents))
                text = "\n".join([para.text for para in doc.paragraphs if para.text])
                return {
                    "success": True,
                    "file_type": "docx",
                    "filename": filename,
                    "text": text.strip(),
                    "paragraphs": len(doc.paragraphs)
                }
            except ImportError:
                raise HTTPException(status_code=500, detail="python-docx is required for DOCX processing")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read DOCX: {str(e)}")

        # Handle TXT files
        elif filename_lower.endswith(".txt"):
            text = contents.decode("utf-8", errors="ignore")
            return {
                "success": True,
                "file_type": "txt",
                "filename": filename,
                "text": text.strip()
            }

        # Handle image files
        elif filename_lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
            try:
                from PIL import Image
                import base64
                image = Image.open(io.BytesIO(contents))
                # Convert to base64 for display
                img_byte_arr = io.BytesIO()
                image.save(img_byte_arr, format='PNG')
                img_byte_arr = img_byte_arr.getvalue()
                img_base64 = base64.b64encode(img_byte_arr).decode('utf-8')
                
                return {
                    "success": True,
                    "file_type": "image",
                    "filename": filename,
                    "image_data": img_base64,
                    "size": f"{image.width}x{image.height}",
                    "format": image.format
                }
            except ImportError:
                raise HTTPException(status_code=500, detail="Pillow is required for image processing")
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to read image: {str(e)}")

        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Please upload PDF, DOCX, TXT, or image files.")

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"File analysis failed: {str(e)}")